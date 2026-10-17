import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# 1. Define the Testplan Data
data = [
    {
        "SNo.": 1,
        "Testname": "TC_BOOT_01",
        "Description": "Power-on Reset & Entry Vector Check: Reset de-assertion allows VeeR-EL2 to fetch its first instruction from address 0x0000_0000 (_start) into ICCM.",
        "Checkpoint": "CPU executes first opcode at 0x0000_0000. Reset values asserted across all peripheral registers.",
        "C verification part": "Execute entry bootstrap in crt0.S, configure gp, set stack pointer (sp = __stack_top), zero out .bss, and write 0x00000001 to GPIO_OUT[31:16] to signal boot start.",
        "SV verification part": "Drive active-low reset rst_l = 0 for 20 cycles, deassert, and monitor trace_rv_i_address_ip === 32'h0000_0000 and trace_rv_i_valid_ip === 1'b1. Verify all internal peripheral buses are quiescent."
    },
    {
        "SNo.": 2,
        "Testname": "TC_BOOT_02",
        "Description": "DCCM Stack Frame & Memory Integrity: Verify byte, half-word, and word reads and writes to local DCCM memory (0x1000_0000) without generating bus traps.",
        "Checkpoint": "Push/pop stack memory returns expected data without bus-error trap or ECC exceptions.",
        "C verification part": "Allocate stack frames across nested function calls (test_uart(), test_timer(), etc.). Push/pop caller/callee registers. Write and verify walking 1s/0s patterns.",
        "SV verification part": "Snoop CPU LSU bus / DCCM interface signals (dccm_ecc_single_error, dccm_ecc_double_error). Assert zero ECC error flags and confirm trace_rv_i_exception_ip === 1'b0."
    },
    {
        "SNo.": 3,
        "Testname": "TC_UART_01",
        "Description": "UART Configuration & Baud Setup: Verify UART Line Control Register (UART_LCR), Baud Divisor (UART_BAUD), and Interrupt Enable (UART_IER) readback.",
        "Checkpoint": "Registers accept configuration and retain values under standard AXI-Lite access.",
        "C verification part": "Write 0x03 (8N1) to UART_LCR (0x4000_000C), divisor 10 to UART_BAUD (0x4000_0008), and 0x01 to UART_IER (0x4000_0004). Read back and assert equality. Branch to fail_trap(3) on mismatch.",
        "SV verification part": "BFM monitor passively snoops AXI interconnect slave 0 (0x4000_0000). Asserts awready/wready handshakes and verifies bresp === 2'b00 (OKAY)."
    },
    {
        "SNo.": 4,
        "Testname": "TC_UART_02",
        "Description": "UART Transmitter Character Serialization: Transmit byte 'K' (0x4B) across serial pin uart_tx_o at configured baud rate (SIM_BAUD_DIV = 10).",
        "Checkpoint": "Serial pin uart_tx_o outputs frame matching 8N1 timing and data 0x4B.",
        "C verification part": "Poll UART_LSR until LSR_THRE (bit 5) is set, then write 0x4B to UART_THR (0x4000_0000). Poll UART_LSR until LSR_TEMT (bit 6) confirms empty transmitter shift register.",
        "SV verification part": "uart_agent monitor detects falling edge (start bit) on uart_tx_o, samples 8 bits sequentially mid-cycle (repeat (BIT_CYCLES)), and checks parity/stop bit. Flags error if captured data != 8'h4B."
    },
    {
        "SNo.": 5,
        "Testname": "TC_UART_03",
        "Description": "UART Receiver Loopback & Data Injection: Inject byte 0xA5 into uart_rx_i and verify reception in Receiver Buffer Register (UART_RBR).",
        "Checkpoint": "UART_RBR reads 0xA5 and uart_read_irq_o asserts.",
        "C verification part": "Wait in polling loop checking UART_LSR[0] (LSR_DATA_RDY). Read UART_RBR (0x4000_0000) and verify (UART_RBR & 0xFF) == 0xA5. Emit step code 0x0006 to GPIO_OUT[31:16].",
        "SV verification part": "On detecting uart_tx_o == 'K', testbench waits 10 bit periods and drives start bit, data bits (8'hA5), and stop bit into uart_rx_i. Assert uart_read_irq_o === 1'b1."
    },
    {
        "SNo.": 6,
        "Testname": "TC_UART_04",
        "Description": "UART FIFO / Buffer Overrun Detection: Force back-to-back unread RX byte injections into uart_rx_i to trigger receiver overrun status.",
        "Checkpoint": "Overrun Error (OE) bit asserted in LSR; system interconnect remains unblocked.",
        "C verification part": "Allow receiver to receive 2+ characters without reading UART_RBR. Read UART_LSR and verify Overrun Error (OE, bit 1) is asserted. Read UART_RBR to flush.",
        "SV verification part": "Driver transmits multiple bytes without pause. Scoreboard checks that uart_read_irq_o stays asserted and the peripheral does not lock the AXI bus interface."
    },
    {
        "SNo.": 7,
        "Testname": "TC_TMR_01",
        "Description": "System Timer Prescaler & Counter Increment: Program prescaler and step rate, enable timer core, and check counter progression.",
        "Checkpoint": "TIMER_VAL_LO increments monotonically at the configured tick frequency.",
        "C verification part": "Write 0x0000_0901 to TIMER_CFG0 (0x4000_110C), set TIMER_CMP to 0xFFFFFFFF, and enable via TIMER_CTRL = 1. Sample t0 = TIMER_VAL_LO, delay 50 NOPs, sample t1, and assert t1 > t0.",
        "SV verification part": "Snoop AXI slave 1. Verify TIMER_CTRL write takes effect and internal counter advances monotonically every 10 system clocks (prescaler = 9 + 1)."
    },
    {
        "SNo.": 8,
        "Testname": "TC_TMR_02",
        "Description": "System Timer Match Interrupt Generation & W1C: Validate compare match trigger between TIMER_VAL_LO and TIMER_CMP_LO.",
        "Checkpoint": "timer_irq_o asserts on match and clears immediately on write-1-to-clear.",
        "C verification part": "Read TIMER_VAL_LO, program TIMER_CMP_LO = t0 + 100, write 0x1 to TIMER_INTR_EN, and poll TIMER_INTR_ST for assertion. Clear interrupt with W1C (TIMER_INTR_ST = 0x1) and write progress 0x0009.",
        "SV verification part": "Monitor timer_irq_o pin. Assert that timer_irq_o transitions 0 -> 1 exactly when counter reaches compare value, and deasserts immediately upon W1C write completion."
    },
    {
        "SNo.": 9,
        "Testname": "TC_AES_01",
        "Description": "AES FIPS-197 128-bit Encryption Vector: Load standard NIST FIPS-197 128-bit Key and Plaintext vectors and execute encryption.",
        "Checkpoint": "Ciphertext matches NIST standard vector (0x3925841d, 0x02dc09fb, 0xdc118597, 0x196a0b32).",
        "C verification part": "Write AES_KEY[0:3] and AES_TXTIN[0:3]. Assert AES_CTRL_REG = 0x01. Poll AES_STATUS until equal to 2 (complete). Read AES_TXTOUT[0:3] and assert match against AES_CT_W[0:3]. Emit step 0x0013.",
        "SV verification part": "Scoreboard latches inputs, computes AES-128 via reference C/DPI model, monitors aes_busy_o asserting during encryption, checks aes_done_o and aes_done_pulse_o pulses, and verifies read-data bus values."
    },
    {
        "SNo.": 10,
        "Testname": "TC_AES_02",
        "Description": "AES Busy Write-Protection Check: Attempt to overwrite Key and Plaintext registers while encryption engine is in BUSY state.",
        "Checkpoint": "Registers are locked while aes_busy_o == 1; ciphertext calculation remains uncorrupted.",
        "C verification part": "Set AES_CTRL_REG = 0x01. Immediately attempt back-to-back writes of dummy data 0xDEADBEEF to AES_KEY0 and AES_TXTIN0 before polling AES_STATUS. Confirm readback after done matches original FIPS vector.",
        "SV verification part": "AXI monitor verifies slave 2 returns either SLVERR or silently drops writes while aes_busy_o === 1'b1, maintaining encryption integrity."
    },
    {
        "SNo.": 11,
        "Testname": "TC_DMA_01",
        "Description": "DMA Autonomous Burst Transfer (SRAM/UART to DCCM): Configure DMA engine to stream data buffer between peripheral space and internal memory without CPU intervention.",
        "Checkpoint": "Target memory buffer matches source data; DMA status indicates completion and triggers interrupt.",
        "C verification part": "Program DMA_DESC0_SRC = UART_BASE, DMA_DESC0_DST = DCCM_BASE, and enable via DMA_CTRL_REG = DMA_CTRL_START | DMA_CTRL_IRQ_EN. Poll DMA_STAT_REG for DMA_STAT_DONE. Emit step 0x0016.",
        "SV verification part": "Passive bus monitor captures master s01 burst requests. Verifies correct incrementing addresses, burst size (axsize), burst length (axlen), and checks assertion of dma_done_o and dma_irq_o."
    },
    {
        "SNo.": 12,
        "Testname": "TC_DMA_02",
        "Description": "DMA Error Abort & Bus Slave Fault Reporting: Force DMA descriptor to access unmapped/illegal address range (e.g., 0x4000_9000).",
        "Checkpoint": "DMA aborts transfer, sets error status bit, and fires error interrupt without bus hang.",
        "C verification part": "Program DMA_DESC0_DST = 0x4000_9000. Trigger DMA start. Poll DMA_STAT_REG and assert DMA_STAT_ERROR (bit 2) is flagged while dma_done is asserted.",
        "SV verification part": "AXI interconnect returns bresp = 2'b10 (SLVERR) to DMA master port. Assert that dma_error_o === 1'b1 and DMA state machine returns cleanly to IDLE."
    },
    {
        "SNo.": 13,
        "Testname": "TC_WDT_01",
        "Description": "Watchdog Software Keepalive (Software Ping): Verify periodic software pings prevent Watchdog timeout expiration.",
        "Checkpoint": "WDT counter is reloaded; wdt_irq_o and wdt_reset_o remain de-asserted.",
        "C verification part": "Load WDT_LOAD = 0x0010_0000, set WDT_IM = 0x01, enable via WDT_CTRL = EN | AUTOKICK. Delay 20 cycles, write 0xABCDEF01 to WDT_PING. Check WDT_MIS & 0x1 == 0.",
        "SV verification part": "Monitor internal down-counter in dut.u_wdt. Assert that writing to WDT_PING immediately reloads counter to WDT_LOAD value and wdt_irq_o === 1'b0."
    },
    {
        "SNo.": 14,
        "Testname": "TC_WDT_02",
        "Description": "Watchdog Hardware AES Auto-Kick Path: Verify autonomous hardware connection where aes_done_pulse_o reloads the WDT counter.",
        "Checkpoint": "WDT reloads synchronously upon aes_done_pulse_o assertion without software intervention.",
        "C verification part": "Enable WDT with auto-kick enabled. Wait for counter to decrement, trigger an AES encryption cycle. Check masked interrupt status WDT_MIS remains 0.",
        "SV verification part": "Monitor the hardware line connecting aes_done_pulse_o to the WDT reload input. Verify counter resets to 0x0010_0000 coincident with the 1-cycle AES pulse."
    },
    {
        "SNo.": 15,
        "Testname": "TC_WDT_03",
        "Description": "Watchdog Early Warning Interrupt & W1C Clear: Allow WDT timer to starve with auto-kick disabled to verify stage 1 early-warning interrupt.",
        "Checkpoint": "wdt_irq_o asserts on counter zero; W1C write clears the pending status.",
        "C verification part": "Program WDT_LOAD = 0x10, disable auto-kick (WDT_CTRL = EN). Poll WDT_RIS until bit 0 is asserted. Clear interrupt via WDT_IC = 0xFFFFFFFF. Confirm WDT_RIS == 0. Emit step 0x0021.",
        "SV verification part": "Assert wdt_irq_o === 1'b1 precisely when counter expires to zero. Check that write to WDT_IC clears wdt_irq_o within 2 clock cycles."
    },
    {
        "SNo.": 16,
        "Testname": "TC_WDT_04",
        "Description": "Watchdog Escalation to NMI & Hard Reset: Verify stage 2 escalation triggers core NMI and asserts system reset line.",
        "Checkpoint": "wdt_nmi_o asserts, followed by wdt_reset_o pulsing active, returning top-level SoC to reset.",
        "C verification part": "Program WDT_CTRL = EN | NMI_EN | RESET_EN and short WDT_LOAD = 0x08. Allow counter to expire without clearing early-warning interrupt.",
        "SV verification part": "Testbench observes wdt_nmi_o === 1'b1 followed by wdt_reset_o === 1'b1. Assert top-level SoC reset resets core registers and returns interconnect to default state."
    },
    {
        "SNo.": 17,
        "Testname": "TC_GPIO_01",
        "Description": "GPIO Pin Direction & Pin Out Drive: Configure pin directions, drive static values, and observe external IO pins.",
        "Checkpoint": "gpio_out_o and gpio_tx_en_o reflect programmed register states exactly.",
        "C verification part": "Write GPIO_EN_REG = 0x0000_00FF to enable lower 8 pins as outputs. Write GPIO_OUT_REG = 0xA5. Read back and confirm output register reflection.",
        "SV verification part": "Sample top-level pins gpio_out_o[7:0] and gpio_tx_en_o[7:0]. Check gpio_out_o[7:0] === 8'hA5 and gpio_tx_en_o[7:0] === 8'hFF."
    },
    {
        "SNo.": 18,
        "Testname": "TC_GPIO_02",
        "Description": "GPIO External Pin Input Sampling: Drive stimulus into gpio_in_i and verify CPU samples the values through synchronizer stages.",
        "Checkpoint": "GPIO_IN_REG returns injected external pattern after 2-clock synchronization.",
        "C verification part": "Execute read of GPIO_IN_REG (0x4000_5100) and compare read data against injected vector (0x5A5AA5A5).",
        "SV verification part": "Drive gpio_in_i = 32'h5A5A_A5A5. Verify signals pass through 2-stage synchronizer (gpio_input_stage) and assert stable read response on interconnect slave 5 within 3 clock cycles."
    },
    {
        "SNo.": 19,
        "Testname": "TC_GPIO_03",
        "Description": "GPIO Edge-Triggered Interrupt Detection: Configure rising-edge trigger on pin 0 and verify interrupt generation on low-to-high transition.",
        "Checkpoint": "gpio_irq_o triggers on 0->1 transition and clears on status register write.",
        "C verification part": "Program GPIO_RISE_EN = 0x0000_0001. Wait for hardware event, then poll GPIO_INT_STAT until bit 0 asserts. Clear interrupt via GPIO_INT_STAT = 0xFFFFFFFF. Emit step 0x0025.",
        "SV verification part": "Wait until CPU writes 0xA5 to gpio_out_o[7:0], then pulse gpio_in_i[0] from 0 -> 1. Assert gpio_irq_o === 1'b1. Verify deassertion after software clears status."
    },
    {
        "SNo.": 20,
        "Testname": "TC_PIPE_01",
        "Description": "Full Pipeline: UART -> DMA -> AES -> WDT: Exercise integrated datalogger loop: byte arrival triggers DMA, feeds AES block, and auto-kicks watchdog.",
        "Checkpoint": "Continuous autonomous pipeline execution without buffer overflow or WDT expiration.",
        "C verification part": "Initialize WDT auto-kick, prime AES key/plaintext, trigger UART reception, and assert no watchdog timeout occurs during continuous data encryption. Emit step 0x0026.",
        "SV verification part": "Inject sequential packets on uart_rx_i. Verify UART interrupt fires, DMA transfers data, AES engine executes, aes_done_pulse_o kicks WDT, and wdt_reset_o remains deasserted throughout."
    },
    {
        "SNo.": 21,
        "Testname": "TC_BUS_01",
        "Description": "Interconnect Multi-Master Arbitration Stress: Force concurrent AXI transactions from CPU LSU Master (s00) and DMA Master (s01) targeting distinct slaves.",
        "Checkpoint": "Zero deadlock; round-robin arbiter cleanly serves both masters without starving either port.",
        "C verification part": "Execute alternating bursts to Timer and AES while DMA is actively reading from UART buffer space.",
        "SV verification part": "BFM driver stresses slave interfaces simultaneously. Verify round-robin arbiter prevents starvation, preserves transaction IDs (arid/awid), and prevents deadlocks."
    },
    {
        "SNo.": 22,
        "Testname": "TC_BUS_02",
        "Description": "Unmapped Address Hole Access (SLVERR Check): Verify bus access to unmapped peripheral space (0x4000_9000 or 0xF000_0000) triggers error response.",
        "Checkpoint": "Interconnect returns SLVERR (bresp/rresp = 2'b10); CPU encounters precise bus-error trap.",
        "C verification part": "Perform 32-bit load/store to address 0x4000_9000. Catch expected bus error trap in mtvec handler.",
        "SV verification part": "Monitor interconnect master bus. Verify slave decoder asserts rresp = 2'b10 or bresp = 2'b10 (SLVERR) without hanging arready/awready."
    },
    {
        "SNo.": 23,
        "Testname": "TC_FULL_01",
        "Description": "Full Regression Firmware Pass Verification: Complete execution of entire firmware suite validating all peripherals sequentially.",
        "Checkpoint": "gpio_out_o[31:16] reports 0x900D; simulation finishes cleanly without timeouts.",
        "C verification part": "Execute main() through test_uart(), test_timer(), test_aes(), test_wdt(), test_gpio(), and test_pipeline(). Write 0x900D to GPIO_OUT[31:16] on final success.",
        "SV verification part": "Testbench samples gpio_out_o[31:16]. Detects 16'h900D, displays '*** ALL FIRMWARE CHECKS PASSED ***', and cleanly calls $finish. Halts with failure report if 16'hDEAD appears."
    }
]

# 2. Convert to DataFrame
df = pd.DataFrame(data)

# 3. Export to Excel with openpyxl styling
filename = "SoC_Datalogger_Verification_Testplan.xlsx"
with pd.ExcelWriter(filename, engine="openpyxl") as writer:
    df.to_excel(writer, sheet_name="SoC_Testplan", index=False)
    worksheet = writer.sheets["SoC_Testplan"]

    # Header styling (Dark Navy with White Bold text)
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="top", wrap_text=True)

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    # Style Header Row
    for col_idx in range(1, len(df.columns) + 1):
        cell = worksheet.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    # Alternating row colors
    fill_even = PatternFill(start_color="F2F5F8", end_color="F2F5F8", fill_type="solid")
    fill_odd = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    # Style Data Rows
    for row_idx in range(2, len(df) + 2):
        row_fill = fill_even if row_idx % 2 == 0 else fill_odd
        for col_idx in range(1, len(df.columns) + 1):
            cell = worksheet.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            cell.fill = row_fill
            if col_idx in [1, 2]:
                cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
                if col_idx == 2:
                    cell.font = Font(name="Calibri", size=10, bold=True)
            else:
                cell.alignment = left_align
                cell.font = Font(name="Calibri", size=10)

    # Set Column Widths
    col_widths = {
        "A": 8,   # SNo.
        "B": 16,  # Testname
        "C": 35,  # Description
        "D": 30,  # Checkpoint
        "E": 45,  # C verification part
        "F": 45   # SV verification part
    }
    for col_letter, width in col_widths.items():
        worksheet.column_dimensions[col_letter].width = width

    # Freeze Header Row
    worksheet.freeze_panes = "A2"

print(f"Generated successfully: {filename}")
